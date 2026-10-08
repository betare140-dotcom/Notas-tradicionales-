"""Lectura de Excel y generación de Word sin servicios externos."""
from datetime import date, datetime
from io import BytesIO
import math
import re
import unicodedata
from urllib.parse import urlsplit

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

FIELDS = ('Fecha', 'Medio', 'Título', 'Link')
SUMMARY_ALIASES = ('resumen', 'extracto', 'sintesis', 'contenido', 'texto', 'contenido de la nota')
ALIASES = {
    'Fecha': ('fecha', 'fecha de nota', 'fecha de la nota', 'fecha publicacion', 'fecha de publicacion', 'published date'),
    'Medio': ('portal', 'medio', 'nombre del medio', 'nombre medio', 'fuente'),
    'Título': ('titulo descriptivo', 'titulo', 'titulo de la nota', 'titulo de nota', 'titular'),
    'Link': ('url', 'link de nota', 'link de la nota', 'link', 'enlace', 'abrir enlace'),
}
MONTHS = ('enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre')

def text(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ''
    # Remove controls forbidden by XML while preserving accents and punctuation.
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value)).strip()

def normalize(value):
    return ''.join(c for c in unicodedata.normalize('NFD', text(value).lower()) if unicodedata.category(c) != 'Mn').strip()

def suggest_mapping(headers):
    normalized = [normalize(h) for h in headers]
    return {field: next((normalized.index(alias) for alias in aliases if alias in normalized), None)
            for field, aliases in ALIASES.items()}

def detect_header(sheet):
    candidates = []
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 50)):
        mapping = suggest_mapping([c.value for c in row])
        candidates.append((sum(v is not None for v in mapping.values()), row[0].row))
    return max(candidates, key=lambda x: (x[0], -x[1]))[1] if candidates else 1

def load_excel(data):
    return load_workbook(BytesIO(data), data_only=True)

def parse_date(value, epoch):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            converted = from_excel(value, epoch)
            return converted.date() if isinstance(converted, datetime) else None
        except (ValueError, OverflowError):
            return None
    raw = text(value)
    try:
        return datetime.fromisoformat(raw.replace('Z', '+00:00')).date()
    except ValueError:
        pass
    for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y', '%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%Y/%m/%d'):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    match = re.fullmatch(r'(\d{1,2}) de ([a-z]+) de (\d{4})', normalize(raw))
    if match and match[2] in MONTHS:
        try:
            return date(int(match[3]), MONTHS.index(match[2]) + 1, int(match[1]))
        except ValueError:
            pass
    return None

def valid_url(value):
    try:
        parsed = urlsplit(value)
        return parsed.scheme.lower() in ('http', 'https') and bool(parsed.netloc) and not any(c.isspace() for c in value)
    except ValueError:
        return False

def read_records(sheet, header_row, mapping, epoch, summary_column=None):
    records = []
    for row in sheet.iter_rows(min_row=header_row + 1):
        cells = {field: row[index] for field, index in mapping.items()}
        if not any(text(c.value) or c.hyperlink for c in cells.values()):
            continue
        link_cell = cells['Link']
        raw_link = text(link_cell.hyperlink.target if link_cell.hyperlink else link_cell.value)
        record = {field: text(c.value) for field, c in cells.items()}
        record['Link'] = raw_link
        record['Resumen'] = text(row[summary_column].value) if summary_column is not None else ''
        record['_date'] = parse_date(cells['Fecha'].value, epoch)
        record['_row'] = row[0].row
        record['_warnings'] = []
        if not record['_date']:
            record['_warnings'].append('Fecha no reconocida')
        for field in ('Medio', 'Título', 'Link'):
            if not record[field]:
                record['_warnings'].append(f'{field} vacío')
        if raw_link and not valid_url(raw_link):
            record['_warnings'].append('Link no válido')
        records.append(record)
    return records

def display_date(record):
    return record['_date'].strftime('%d/%m/%Y') if record['_date'] else (record['Fecha'] or 'Sin fecha')

def date_long(d):
    return f'{d.day} de {MONTHS[d.month-1]} de {d.year}'

def hyperlink(paragraph, label, url, line_break=False):
    element = OxmlElement('w:hyperlink')
    element.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1')
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single')
    props.extend([color, underline]); run.append(props)
    if line_break:
        run.append(OxmlElement('w:br'))
    content = OxmlElement('w:t'); content.text = label
    run.append(content); element.append(run); paragraph._p.append(element)

def create_word(records, title='Notas de medios tradicionales'):
    doc = Document()
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.8)
    for style in doc.styles:
        if hasattr(style, 'font'):
            style.font.name = 'Verdana'
            style.font.size = Pt(12)
    for name in ('Normal', 'Title', 'Heading 1'):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles['Normal']
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(9)
    doc.styles['Title'].font.size = Pt(12)
    doc.add_paragraph(text(title), 'Title')
    doc.add_paragraph(f'{len(records)} notas de medios tradicionales.')
    groups = {}
    for record in records:
        key = record['_date'] or display_date(record)
        groups.setdefault(key, []).append(record)
    for grouped in groups.values():
        heading = doc.add_paragraph()
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_before = Pt(10)
        heading.paragraph_format.space_after = Pt(6)
        run = heading.add_run(date_long(grouped[0]['_date']) if grouped[0]['_date'] else display_date(grouped[0]))
        run.bold = True
        run.font.color.rgb = RGBColor.from_string('0563C1')
        run.font.size = Pt(12)
        for record in grouped:
            add_note(doc, record)
    # Include OOXML hyperlink runs, which python-docx paragraph.runs omits.
    for run in doc.element.iter(qn('w:r')):
        props = run.find(qn('w:rPr'))
        if props is None:
            props = OxmlElement('w:rPr'); run.insert(0, props)
        fonts = props.find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts'); props.insert(0, fonts)
        for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn(f'w:{attr}'), 'Verdana')
        for tag in ('sz', 'szCs'):
            size = props.find(qn(f'w:{tag}'))
            if size is None:
                size = OxmlElement(f'w:{tag}'); props.append(size)
            size.set(qn('w:val'), '24')
    stream = BytesIO(); doc.save(stream)
    return stream.getvalue()

def add_note(doc, record):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    for i, (field, fallback) in enumerate((('Medio', 'Sin medio'), ('Título', 'Sin título'), ('Resumen', 'Extracto pendiente de completar.'))):
        run = p.add_run()
        if i:
            run.add_break()
        run.add_text(text(record.get(field)) or fallback)
        run.bold = field == 'Medio'
    if valid_url(record['Link']):
        hyperlink(p, record['Link'], record['Link'], line_break=True)
    else:
        p.add_run().add_break()
        p.add_run(record['Link'] or 'Sin enlace')
